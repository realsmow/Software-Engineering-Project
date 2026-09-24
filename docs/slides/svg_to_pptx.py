"""Convert a folder of slide SVGs into one editable .pptx for Canva.

    cd docs/slides
    .venv/bin/python svg_to_pptx.py demo-slides-02            # -> demo-slides-02/demo-slides-02.pptx
    .venv/bin/python svg_to_pptx.py demo-slides-02 -o out.pptx

Why not just drop the PNGs into Canva — a full-slide picture cannot be edited.
Canva's PPTX import keeps every shape as its own element instead, so after
this conversion each text, box, circle, line and screenshot can be moved,
retyped or recoloured in Canva.

Only the SVG subset our generators emit is supported: <text>, <rect>,
<circle>, <line>, <path>/<polyline> made of straight segments, <image>
(optionally wrapped in a nested <svg> whose viewBox crops it) and
<g transform="translate(...)">. Anything else is reported, not guessed at.

Needs python-pptx and fonttools, installed in docs/slides/.venv (see README).
"""

from __future__ import annotations

import argparse
import base64
import io
import re
import subprocess
import sys
from functools import lru_cache
from pathlib import Path

from fontTools.ttLib import TTFont
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu

SVG_NS = "{http://www.w3.org/2000/svg}"
PX = 9525            # EMU per CSS px at 96 dpi, so a 1920x1080 SVG becomes a 20x11.25 in slide
PT_PER_PX = 0.75

# SVG font-family (first name in the list) -> font name written into the PPTX.
# The targets must exist in Canva, otherwise Canva substitutes a random font.
# DejaVu Sans Mono is not in Canva; Roboto Mono has the same 0.6em advance.
FONT_MAP = {
    "Sarabun": "Sarabun",
    "Laksaman": "Sarabun",        # Laksaman is a TH Sarabun clone, Sarabun is the closest Canva font
    "Noto Sans Thai": "Noto Sans Thai",
    "DejaVu Sans Mono": "Roboto Mono",
}


# ------------------------------------------------------------------ fonts
@lru_cache(maxsize=None)
def _font_file(family: str, bold: bool) -> str:
    pattern = family + (":bold" if bold else "")
    return subprocess.run(["fc-match", "-f", "%{file}", pattern],
                          check=True, capture_output=True, text=True).stdout.strip()


@lru_cache(maxsize=None)
def _measure_font(family: str, bold: bool) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(_font_file(family, bold), 100)


@lru_cache(maxsize=None)
def _baseline_offset(family: str) -> float:
    """Distance (in em) from a centred line's middle down to its baseline.

    Renderers place text differently: PowerPoint uses the Windows metrics,
    LibreOffice and browsers use hhea. Anchoring each box at its vertical
    middle makes that difference mostly cancel, because every renderer
    centres the line box it built. With hhea ascent a and descent d the
    baseline sits (a - d) / 2 below the middle.
    """
    font = TTFont(_font_file(family, False), lazy=True)
    upm = font["head"].unitsPerEm
    hhea = font["hhea"]
    return (hhea.ascent + hhea.descent) / 2 / upm   # hhea.descent is negative


def text_width(s: str, size: float, family: str, bold: bool) -> float:
    return _measure_font(family, bold).getlength(s) * size / 100


# ------------------------------------------------------------------ small parsers
def num(el, name: str, default: float = 0.0) -> float:
    v = el.get(name)
    return float(v) if v is not None else default


def color(value: str | None) -> RGBColor | None:
    if not value or value == "none":
        return None
    value = value.lstrip("#")
    if len(value) == 3:
        value = "".join(c * 2 for c in value)
    return RGBColor.from_string(value.upper())


def opacity(el, *names: str) -> float:
    """Multiply opacity and any of the given *-opacity attributes."""
    a = num(el, "opacity", 1.0)
    for n in names:
        a *= num(el, n, 1.0)
    return a


def set_alpha(fill_or_line_fill, alpha: float) -> None:
    """python-pptx has no alpha API; add <a:alpha> under the solid colour."""
    if alpha >= 0.999:
        return
    clr = fill_or_line_fill._xPr.find(".//" + qn("a:srgbClr"))
    if clr is None:
        return
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def drop_theme_style(shape) -> None:
    """Remove the <p:style> python-pptx gives every autoshape and connector.

    It points at the theme's accent colour and effect style, and some renderers
    (LibreOffice, and Canva's importer) draw that effect as a drop shadow even
    when fill and line are set explicitly. Our shapes carry their own fill and
    line, so the theme reference has nothing left to contribute.
    """
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def parse_translate(transform: str | None) -> tuple[float, float]:
    if not transform:
        return 0.0, 0.0
    m = re.fullmatch(r"\s*translate\(\s*([-\d.]+)[\s,]*([-\d.]*)\s*\)\s*", transform)
    if not m:
        raise ValueError(f"unsupported transform: {transform}")
    return float(m.group(1)), float(m.group(2) or 0)


def path_contours(d: str) -> list[list[tuple[float, float]]]:
    """Straight-line SVG path data -> list of point lists (one per subpath).

    Supports M m L l H h V v Z z, which is all our generators write.
    """
    tokens = re.findall(r"[MmLlHhVvZz]|-?\d*\.?\d+(?:e-?\d+)?", d)
    contours: list[list[tuple[float, float]]] = []
    x = y = 0.0
    cmd = None
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.isalpha():
            cmd = t
            i += 1
            if cmd in "Zz":
                if contours and contours[-1]:
                    contours[-1].append(contours[-1][0])
                continue
        if cmd is None:
            raise ValueError(f"path data does not start with a command: {d}")
        if cmd in "MmLl":
            dx, dy = float(tokens[i]), float(tokens[i + 1])
            i += 2
            rel = cmd.islower()
            x, y = (x + dx, y + dy) if rel else (dx, dy)
            if cmd in "Mm":
                contours.append([(x, y)])
                cmd = "l" if rel else "L"     # extra pairs after M are line-tos
            else:
                contours[-1].append((x, y))
        elif cmd in "HhVv":
            v = float(tokens[i])
            i += 1
            if cmd == "H":
                x = v
            elif cmd == "h":
                x += v
            elif cmd == "V":
                y = v
            else:
                y += v
            contours[-1].append((x, y))
        else:
            raise ValueError(f"unsupported path command {cmd!r} in {d}")
    return contours


# ------------------------------------------------------------------ converter
class SlideWriter:
    """Walks one SVG tree and adds matching shapes to one slide, in paint order."""

    def __init__(self, slide):
        self.shapes = slide.shapes
        self.skipped: dict[str, int] = {}

    # --- geometry helper: SVG px -> EMU
    @staticmethod
    def e(v: float) -> Emu:
        return Emu(round(v * PX))

    def walk(self, el, ox: float = 0.0, oy: float = 0.0) -> None:
        for child in el:
            if not isinstance(child.tag, str):
                continue          # comments
            tag = child.tag.replace(SVG_NS, "")
            handler = getattr(self, f"do_{tag}", None)
            if handler is None:
                self.skipped[tag] = self.skipped.get(tag, 0) + 1
                continue
            handler(child, ox, oy)

    # --- containers
    def do_g(self, el, ox, oy):
        dx, dy = parse_translate(el.get("transform"))
        self.walk(el, ox + dx, oy + dy)

    def do_svg(self, el, ox, oy):
        """Nested <svg> = a cropped screenshot. viewBox picks the crop window."""
        x, y = ox + num(el, "x"), oy + num(el, "y")
        w, h = num(el, "width"), num(el, "height")
        vb = [float(v) for v in el.get("viewBox", f"0 0 {w} {h}").split()]
        for img in el.iter(SVG_NS + "image"):
            self.add_image(img, x, y, w, h, vb)

    # --- leaves
    def do_image(self, el, ox, oy):
        x, y = ox + num(el, "x"), oy + num(el, "y")
        w, h = num(el, "width"), num(el, "height")
        self.add_image(el, x, y, w, h, None)

    def add_image(self, img_el, x, y, w, h, viewbox):
        href = img_el.get("href") or img_el.get("{http://www.w3.org/1999/xlink}href")
        if not href or not href.startswith("data:"):
            self.skipped["image(external)"] = self.skipped.get("image(external)", 0) + 1
            return
        raw = base64.b64decode(href.split(",", 1)[1])
        im = Image.open(io.BytesIO(raw))
        if viewbox is not None:
            # Crop in Pillow instead of using PPTX's srcRect cropping: every
            # importer honours a plain picture, not every one honours srcRect.
            # The image inside is placed at 0,0 at its natural size, so the
            # viewBox is directly a pixel box on the source image.
            vx, vy, vw, vh = viewbox
            im = im.crop((round(vx), round(vy), round(vx + vw), round(vy + vh)))
        buf = io.BytesIO()
        im.convert("RGB").save(buf, "JPEG", quality=92)
        buf.seek(0)
        self.shapes.add_picture(buf, self.e(x), self.e(y), self.e(w), self.e(h))

    def style_fill_line(self, shape, el):
        fill = color(el.get("fill", "#000"))
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
            set_alpha(shape.fill, opacity(el, "fill-opacity"))
        stroke = color(el.get("stroke"))
        if stroke is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = stroke
            shape.line.width = self.e(num(el, "stroke-width", 1.0))
            set_alpha(shape.line.fill, opacity(el, "stroke-opacity"))
        drop_theme_style(shape)

    def do_rect(self, el, ox, oy):
        x, y = ox + num(el, "x"), oy + num(el, "y")
        w, h = num(el, "width"), num(el, "height")
        rx = num(el, "rx", num(el, "ry"))
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0 else MSO_SHAPE.RECTANGLE
        shape = self.shapes.add_shape(kind, self.e(x), self.e(y), self.e(w), self.e(h))
        if rx > 0:
            # PPTX corner radius is a fraction of the shorter side, capped at 0.5
            shape.adjustments[0] = min(0.5, rx / min(w, h))
        self.style_fill_line(shape, el)

    def do_circle(self, el, ox, oy):
        cx, cy, r = ox + num(el, "cx"), oy + num(el, "cy"), num(el, "r")
        shape = self.shapes.add_shape(MSO_SHAPE.OVAL, self.e(cx - r), self.e(cy - r),
                                      self.e(2 * r), self.e(2 * r))
        self.style_fill_line(shape, el)

    def do_line(self, el, ox, oy):
        conn = self.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            self.e(ox + num(el, "x1")), self.e(oy + num(el, "y1")),
            self.e(ox + num(el, "x2")), self.e(oy + num(el, "y2")))
        conn.line.color.rgb = color(el.get("stroke", "#000"))
        conn.line.width = self.e(num(el, "stroke-width", 1.0))
        set_alpha(conn.line.fill, opacity(el, "stroke-opacity"))
        drop_theme_style(conn)

    def do_path(self, el, ox, oy):
        self.add_polyline(path_contours(el.get("d", "")), el, ox, oy)

    def do_polyline(self, el, ox, oy):
        nums = [float(v) for v in re.split(r"[\s,]+", el.get("points", "").strip()) if v]
        self.add_polyline([list(zip(nums[::2], nums[1::2]))], el, ox, oy)

    def add_polyline(self, contours, el, ox, oy):
        contours = [c for c in contours if len(c) > 1]
        if not contours:
            return
        (sx, sy), *rest = contours[0]
        fb = self.shapes.build_freeform(self.e(ox + sx), self.e(oy + sy), scale=1.0)
        fb.add_line_segments([(self.e(ox + x), self.e(oy + y)) for x, y in rest], close=False)
        for c in contours[1:]:
            fb.move_to(self.e(ox + c[0][0]), self.e(oy + c[0][1]))
            fb.add_line_segments([(self.e(ox + x), self.e(oy + y)) for x, y in c[1:]],
                                 close=False)
        shape = fb.convert_to_shape()
        self.style_fill_line(shape, el)

    def do_text(self, el, ox, oy):
        s = "".join(el.itertext())
        if not s.strip():
            return
        family = el.get("font-family", "sans-serif").split(",")[0].strip().strip("'\"")
        size = num(el, "font-size", 16)
        bold = int(el.get("font-weight", "400").replace("bold", "700")) >= 600
        anchor = el.get("text-anchor", "start")
        spacing = num(el, "letter-spacing")
        x, y = ox + num(el, "x"), oy + num(el, "y")

        # Pad the box generously. Canva re-wraps text to its box width, so a box
        # that is exactly as wide as the text wraps the moment Canva's shaping
        # comes out a pixel wider than ours. The padding is invisible.
        w = text_width(s, size, family, bold) + spacing * len(s)
        box_w = w * 1.12 + size
        left = {"start": x, "middle": x - box_w / 2, "end": x - box_w}[anchor]
        box_h = size * 1.6
        mid = y - _baseline_offset(family) * size

        tb = self.shapes.add_textbox(self.e(left), self.e(mid - box_h / 2),
                                     self.e(box_w), self.e(box_h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = None
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER,
                       "end": PP_ALIGN.RIGHT}[anchor]
        run = p.add_run()
        run.text = s
        font = run.font
        font.size = Emu(round(size * PT_PER_PX * 12700))
        font.bold = bold
        font.color.rgb = color(el.get("fill", "#000"))
        set_alpha(run.font.fill, opacity(el, "fill-opacity"))

        # Thai is a complex script: PowerPoint and Canva read the font for Thai
        # characters from <a:cs>, not <a:latin>. Setting only font.name would
        # leave all the Thai text in the theme's default font.
        name = FONT_MAP.get(family, family)
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = etree.SubElement(rpr, qn(tag))
            node.set("typeface", name)
        rpr.set("lang", "th-TH")
        if spacing:
            rpr.set("spc", str(round(spacing * PT_PER_PX * 100)))   # hundredths of a point


def _order_rpr_children(prs_root) -> None:
    """Keep <a:rPr> children in schema order (fill before latin/ea/cs).

    SubElement appends at the end, and python-pptx adds <a:solidFill> when the
    colour is set — the order they end up in depends on call order. Strict
    readers (PowerPoint, Canva) reject out-of-order children, so fix it once.
    """
    order = ["a:ln", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill",
             "a:grpFill", "a:effectLst", "a:effectDag", "a:highlight", "a:uLnTx", "a:uLn",
             "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick",
             "a:hlinkMouseOver", "a:rtl", "a:extLst"]
    rank = {qn(t): i for i, t in enumerate(order)}
    for rpr in prs_root.iter(qn("a:rPr")):
        kids = sorted(rpr, key=lambda k: rank.get(k.tag, len(order)))
        for k in kids:
            rpr.append(k)       # re-appending moves the node, giving sorted order


def convert(folder: Path, out: Path) -> None:
    svgs = sorted(folder.glob("*.svg"))
    if not svgs:
        raise SystemExit(f"no .svg files in {folder}")
    first = etree.parse(str(svgs[0])).getroot()
    prs = Presentation()
    prs.slide_width = Emu(round(num(first, "width") * PX))
    prs.slide_height = Emu(round(num(first, "height") * PX))
    blank = prs.slide_layouts[6]

    for path in svgs:
        root = etree.parse(str(path)).getroot()
        slide = prs.slides.add_slide(blank)
        writer = SlideWriter(slide)
        writer.walk(root)
        _order_rpr_children(slide.shapes._spTree)
        note = f"  ข้าม {writer.skipped}" if writer.skipped else ""
        print(f"  {path.name}  {len(slide.shapes)} ชิ้น{note}")

    prs.save(str(out))
    print(f"บันทึก {out}  ({out.stat().st_size / 1e6:.1f} MB, {len(svgs)} สไลด์)")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("folder", type=Path, help="folder of NN-name.svg slides")
    ap.add_argument("-o", "--out", type=Path, help="output .pptx (default: inside the folder)")
    args = ap.parse_args()
    out = args.out or args.folder / f"{args.folder.name}.pptx"
    convert(args.folder, out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
